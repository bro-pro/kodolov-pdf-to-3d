# -*- coding: utf-8 -*-
"""Экспорт Project Model JSON в Excel для тестового задания Kodolov."""
import json
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

if len(sys.argv) < 2:
    print("Использование: python export_excel_p2_6.py model.json [estimate.xlsx]")
    raise SystemExit(2)

src = Path(sys.argv[1])
out = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_name("kodolov_estimate.xlsx")
d = json.loads(src.read_text(encoding="utf-8"))

rooms = d.get("rooms", [])
walls = {w["id"]: w for w in d.get("walls", [])}
openings = {o["id"]: o for o in d.get("openings", [])}
q = d.get("quantities", {})

def room_name(r):
    ref = r.get("reference_label_area_m2")
    return f"Помещение {r['id'].split('_')[-1]} ({ref:.2f} м² по плану)" if ref else f"Помещение {r['id'].split('_')[-1]}"

wb = Workbook()
ws = wb.active
ws.title = "Объёмы по помещениям"

headers = [
    "Помещение", "ID", "Площадь пола, м²", "Площадь стен gross, м²",
    "Площадь стен net, м²", "Площадь потолка, м²", "Периметр, м",
    "Плинтус, м", "Двери, шт.", "Окна, шт.",
    "Площадь дверных проёмов, м²", "Площадь оконных проёмов, м²", "Статус"
]
ws.append(headers)

for r in rooms:
    wall_ids = list(dict.fromkeys(r.get("wall_ids", [])))
    door_ids = [x for x in dict.fromkeys(r.get("door_ids", []))
                if x in openings and openings[x].get("kind") == "door"]
    win_ids = [x for x in dict.fromkeys(r.get("window_ids", []))
               if x in openings and openings[x].get("kind") == "window"]
    gross = sum(walls[wid].get("gross_area_one_side_m2", 0) for wid in wall_ids if wid in walls)
    door_area = sum(openings[oid].get("width_m", 0) * openings[oid].get("height_m", 2.1) for oid in door_ids)
    win_area = sum(
        openings[oid].get("width_m", 0) *
        (openings[oid].get("top_m", 2.3) - openings[oid].get("bottom_m", 0.9))
        for oid in win_ids
    )
    net = max(0, gross - door_area - win_area)
    baseboard = max(0, r.get("perimeter_m", 0) - sum(openings[oid].get("width_m", 0) for oid in door_ids))
    ws.append([
        room_name(r), r["id"], r.get("area_m2"), gross, net, r.get("area_m2"),
        r.get("perimeter_m"), baseboard, len(door_ids), len(win_ids),
        door_area, win_area, "geometric / topology"
    ])

ws2 = wb.create_sheet("Сводка")
ws2.append(["Показатель", "Ед. изм.", "Количество", "Статус"])
summary = [
    ("Площадь пола", "м²", q.get("floor_area_m2"), q.get("floor_area_status")),
    ("Площадь потолка", "м²", q.get("ceiling_area_m2"), q.get("ceiling_area_status")),
    ("Оси стен", "м", q.get("wall_axis_length_m"), "geometric"),
    ("Площадь стен gross, 1 сторона", "м²", q.get("gross_wall_area_one_side_m2"), "geometric"),
    ("Площадь стен net, 1 сторона", "м²", q.get("net_wall_area_one_side_est_m2"), "estimated"),
    ("Отделка стен, обе стороны", "м²", q.get("wall_finish_area_both_sides_est_m2"), q.get("wall_finish_both_sides_status")),
    ("Объём стен", "м³", q.get("wall_volume_m3"), "geometric"),
    ("Плинтус", "м", q.get("baseboard_length_m"), q.get("baseboard_status")),
    ("Двери", "шт.", q.get("doors_count"), "measured"),
    ("Окна", "шт.", q.get("windows_count"), "mixed"),
    ("Площадь дверных проёмов", "м²", q.get("door_opening_area_m2"), "measured"),
    ("Площадь оконных проёмов", "м²", q.get("window_opening_area_m2"), "estimated"),
]
for row in summary:
    ws2.append(row)

ws3 = wb.create_sheet("Проёмы")
ws3.append(["ID", "Тип", "Источник", "Ширина, м", "Высота, м", "Комнаты", "Стена", "Статус", "Уверенность"])
for o in openings.values():
    h = o.get("height_m", 2.1)
    if o.get("kind") == "window":
        h = o.get("top_m", 2.3) - o.get("bottom_m", 0.9)
    ws3.append([
        o["id"], o.get("kind"), o.get("source"), o.get("width_m"), h,
        ", ".join(o.get("room_ids", [])), o.get("wall_id"),
        o.get("source_status"), o.get("confidence")
    ])

ws4 = wb.create_sheet("Мебель")
ws4.append(["ID", "Тип", "Помещение", "Ширина, м", "Глубина, м", "Высота, м", "X, м", "Y, м", "Поворот, °", "Источник"])
for f in d.get("furniture", {}).get("objects", []):
    ws4.append([
        f.get("id"), f.get("type"), f.get("room_id"), f.get("width_m"),
        f.get("depth_m"), f.get("height_m"), f.get("x_m"), f.get("y_m"),
        f.get("rotation_deg"), f.get("source")
    ])

ws5 = wb.create_sheet("Контроль")
checks = [
    ("Исходный PDF", d.get("source_pdf")),
    ("Лист PDF", d.get("page")),
    ("Схема модели", d.get("schema")),
    ("Источник масштаба", d.get("scale", {}).get("source")),
    ("Масштаб, м/точку", d.get("scale", {}).get("m_per_pdf_point")),
    ("Уверенность масштаба", d.get("scale", {}).get("confidence")),
    ("Комнаты", q.get("room_count")),
    ("Семантические названия придуманы", d.get("validation", {}).get("semantic_room_names_invented")),
    ("Room↔wall проверено", d.get("validation", {}).get("room_wall_links_checked")),
    ("Wall↔opening проверено", d.get("validation", {}).get("wall_opening_links_checked")),
    ("Мебель", d.get("furniture", {}).get("status", "нет")),
    ("Количество мебели", d.get("furniture", {}).get("count", 0)),
]
for a, b in checks:
    ws5.append([a, b])

for sheet in wb.worksheets:
    sheet.freeze_panes = "A2"
    for c in sheet[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows():
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(c.value, float):
                c.number_format = "0.00"
    for col in range(1, sheet.max_column + 1):
        width = max(len(str(sheet.cell(r, col).value or "")) for r in range(1, sheet.max_row + 1))
        sheet.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 12), 42)

wb.save(out)
print(f"Excel сохранён: {out}")
